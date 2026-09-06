"""Generate one-page RIS AI Assist implementation report as DOCX."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from pathlib import Path


def add_heading(doc, text, size=14):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        run.font.size = Pt(10)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    for run in p.runs:
        run.font.size = Pt(10)
    return p


def main():
    out_path = Path(__file__).resolve().parent / "RIS_AI_Assist_Implementation_Report.docx"

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("RIS AI Assist — Implementation Report")
    tr.bold = True
    tr.font.size = Pt(16)
    title.paragraph_format.space_after = Pt(2)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("HGH HisOrder Radiology Information System (Interpretation Desktop)")
    sr.font.size = Pt(10)
    sr.italic = True
    subtitle.paragraph_format.space_after = Pt(10)

    add_heading(doc, "1. Purpose and Scope", 11)
    add_body(
        doc,
        "The RIS AI Assist feature accelerates radiology report drafting on the Interpretation Desktop. "
        "When a radiologist opens an exam, they can request an AI-generated preliminary draft that "
        "populates the Impression field and the CKEditor narrative body. The output is explicitly a "
        "worksheet for licensed radiologist review—not a final diagnosis. The feature integrates the "
        "existing Orthanc PACS archive with OpenAI Vision (GPT-4o) through the hospital MVC application."
    )

    add_heading(doc, "2. End-to-End Workflow", 11)
    add_bullet(doc, "Radiologist opens Interpretation Desktop for an exam with a valid StudyInstanceUID.")
    add_bullet(doc, "User clicks the AI ASSIST button; the browser POSTs to /Radiology/Interpretation/AIAssist.")
    add_bullet(doc, "Server resolves DICOM instances from Orthanc via DICOMweb, renders up to four JPEG previews per study.")
    add_bullet(doc, "OpenAI Vision analyzes de-identified images with modality, age group, and sex context.")
    add_bullet(doc, "JSON response fills #reportImpression (plain text) and CKEditor narrative (HTML); radiologist edits, saves, and finalizes.")

    add_heading(doc, "3. Architecture and Key Components", 11)
    add_body(
        doc,
        "Three image-access paths coexist: OHIF viewer (browser → Orthanc), AI Assist (MVC server → Orthanc via "
        "OrthancInternal HttpClient with configured credentials), and Stream proxy (browser → MVC → Orthanc). "
        "Server-side AI Assist does not rely on browser-cached Orthanc authentication."
    )
    add_bullet(doc, "Backend: AIAssistService.cs, OrthancAccessHelper.cs, InterpretationController.cs")
    add_bullet(doc, "Frontend: Desktop.cshtml, radiology-interpretation.js (toast feedback, field population)")
    add_bullet(doc, "Models/DTOs: AIAssistOptions, AIAssistResult, AIAssistRequestDto, AIAssistUnavailableException")
    add_bullet(doc, "DI registration in Program.cs: IAIAssistService, OrthancInternal and OpenAI HttpClients")

    add_heading(doc, "4. Configuration and Operations", 11)
    add_bullet(doc, "OrthancSettings: ApiBaseUrl, HttpUsername, HttpPassword (must match orthanc.json RegisteredUsers).")
    add_bullet(doc, "AIAssist: OpenAIApiKey, Model (gpt-4o), MaxImages (4), RetryMaxImages (2), ImageDetail (low), MaxTokens (2000).")
    add_bullet(doc, "Environment overrides: ORTHANC_HTTP_USERNAME/PASSWORD, AIAssist__OpenAIApiKey, OPENAI_API_KEY.")
    add_bullet(doc, "OpenAI keys belong in user secrets or local Development config—never in committed source.")

    add_heading(doc, "5. Safety, Resilience, and Clinical Governance", 11)
    add_body(
        doc,
        "Prompts frame the model as a radiology documentation assistant for de-identified PACS renders, "
        "requiring neutral observational language, uncertainty qualifiers, and JSON output "
        "(impression + narrative). Image detail is set to low to reduce policy refusals on medical imaging. "
        "If the primary request fails (refusal or empty response), the service retries with a simplified prompt "
        "and fewer images; partial text may be used as a disclaimer fallback. All successful responses include: "
        "\"AI-generated draft for radiologist review only. Not a final diagnosis.\" HTTP 503 is returned when "
        "Orthanc has no renderable instances; 502/401 scenarios are logged with diagnostic detail for operators."
    )

    add_heading(doc, "6. Summary", 11)
    add_body(
        doc,
        "RIS AI Assist delivers a production-ready, server-mediated bridge between Orthanc DICOM storage and "
        "OpenAI Vision, embedded in the existing interpretation workflow. It reduces documentation time while "
        "preserving radiologist authority: drafts are editable, saveable, and subject to VERIFY & FINALIZE before "
        "becoming an official report."
    )

    doc.save(out_path)
    print(f"Created: {out_path}")


if __name__ == "__main__":
    main()
